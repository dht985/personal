"""下载《云南大学本科学生综合素质测评表》并读取其中与学科竞赛相关的赋分规定。"""

import sys
from pathlib import Path

import requests

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
RAW.mkdir(parents=True, exist_ok=True)
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

H = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Referer": "http://www.xsjzwh.ynu.edu.cn/info/1019/2264.htm",
}

URL = (
    "http://www.xsjzwh.ynu.edu.cn/system/_content/download.jsp"
    "?urltype=news.DownloadAttachUrl&owner=1344356920&wbfileid=6461871"
)

if __name__ == "__main__":
    r = requests.get(URL, headers=H, timeout=40)
    print("status", r.status_code, "bytes", len(r.content))
    out = RAW / "szcp_table.docx"
    out.write_bytes(r.content)
    print("saved ->", out)

    from docx import Document

    doc = Document(str(out))
    print("\n=== 段落 ===")
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            print(" ", text[:160])
    print("\n=== 表格 ===")
    for ti, table in enumerate(doc.tables):
        print(f"--- 表 {ti + 1}（{len(table.rows)} 行 x {len(table.columns)} 列）")
        for row in table.rows[:40]:
            cells = [" ".join(c.text.split())[:60] for c in row.cells]
            line = " | ".join(cells)
            if line.strip(" |"):
                print("   ", line[:220])
